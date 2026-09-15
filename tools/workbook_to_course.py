from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "resources" / "Module_Resource_Collection.xlsx"
OUTPUT = ROOT / "content" / "course.json"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def split_values(value) -> list[str]:
    return [item.strip() for item in re.split(r"[;\n,]+", clean(value)) if item.strip()]


def split_paragraphs(value) -> list[str]:
    text = clean(value)
    if not text:
        return []
    return [item.strip() for item in re.split(r"\n\s*\n|\n", text) if item.strip()]


def step_suffix(value) -> str:
    match = re.search(r"[234]", clean(value))
    return match.group(0) if match else ""


def sheet_rows(workbook, name: str) -> list[dict]:
    sheet = workbook[name]
    headers = [clean(cell.value) for cell in sheet[4]]
    rows = []
    for values in sheet.iter_rows(min_row=5, values_only=True):
        record = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
        if any(clean(value) for value in record.values()):
            rows.append(record)
    return rows


def parse_vocabulary(value) -> list[dict[str, str]]:
    items = []
    for entry in re.split(r"[;\n]+", clean(value)):
        if not entry.strip():
            continue
        term, separator, definition = entry.partition(":")
        items.append({"term": term.strip(), "definition": definition.strip() if separator else "Definition requires designer approval."})
    return items


def choices(row: dict) -> list[dict[str, str]]:
    result = []
    for letter in "ABCDEF":
        text = clean(row.get(f"Option {letter}"))
        if text:
            result.append({"id": letter.lower(), "text": text})
    return result


def question(row: dict) -> dict:
    return {
        "id": clean(row.get("Question ID")),
        "type": clean(row.get("Question Type")).lower().replace("/", "_").replace(" ", "_"),
        "prompt": clean(row.get("Prompt")),
        "choices": choices(row),
        "correct": [item.lower() for item in split_values(row.get("Correct Answer(s)"))],
        "feedback_correct": clean(row.get("Correct Feedback")),
        "feedback_incorrect": clean(row.get("Incorrect Feedback")),
        "source_ids": split_values(row.get("Source IDs")),
        "standard_ids": split_values(row.get("Standard IDs")),
        "objective_ids": split_values(row.get("Objective IDs")),
    }


def main() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK}")
    wb = load_workbook(WORKBOOK, data_only=True)

    setup = {clean(row.get("Setting")): row.get("Value") for row in sheet_rows(wb, "Course Setup")}
    standards = []
    for row in sheet_rows(wb, "Standards"):
        if not clean(row.get("Standard ID")):
            continue
        standards.append({
            "id": clean(row.get("Standard ID")),
            "text": clean(row.get("Full Standard Wording")),
            "organization": clean(row.get("Issuing Organization")),
            "source_url": clean(row.get("Source URL")),
            "priority": clean(row.get("Priority")),
            "notes": clean(row.get("Designer Notes")),
        })

    resources = []
    for row in sheet_rows(wb, "Resources"):
        if clean(row.get("Approval Status")) != "Approved":
            continue
        resources.append({
            "id": clean(row.get("Resource ID")),
            "title": clean(row.get("Title")),
            "type": clean(row.get("Resource Type")),
            "url": clean(row.get("URL or File Path")),
            "publisher": clean(row.get("Publisher or Author")),
            "runtime_minutes": row.get("Runtime Minutes") or 0,
            "transcript_url": clean(row.get("Transcript URL or Path")),
            "description": clean(row.get("Instructional Purpose")),
            "attribution": clean(row.get("Source or Attribution")),
        })

    content_rows = {}
    for row in sheet_rows(wb, "Lesson Content"):
        if clean(row.get("Approval Status")) != "Approved":
            continue
        key = (int(row.get("Lesson")), step_suffix(row.get("Step")))
        content_rows[key] = row

    check_rows = defaultdict(list)
    for row in sheet_rows(wb, "Checks for Understanding"):
        if clean(row.get("Approval Status")) != "Approved":
            continue
        key = (int(row.get("Lesson")), step_suffix(row.get("Step")))
        check_rows[key].append(question(row))

    app_rows = {}
    for row in sheet_rows(wb, "Application Activities"):
        if clean(row.get("Approval Status")) == "Approved":
            app_rows[int(row.get("Lesson"))] = row

    scenarios = defaultdict(list)
    for row in sheet_rows(wb, "Program Scenarios"):
        if clean(row.get("Approval Status")) != "Approved":
            continue
        scenarios[int(row.get("Lesson"))].append({
            "program": clean(row.get("Program")),
            "id": clean(row.get("Scenario ID")),
            "situation": clean(row.get("Situation")),
            "prompts": [clean(row.get(f"Prompt {number}")) for number in range(1, 4) if clean(row.get(f"Prompt {number}"))],
            "exemplar": clean(row.get("Exemplar or Feedback")),
            "source_ids": split_values(row.get("Source IDs")),
        })

    lessons = []
    blueprint_rows = sheet_rows(wb, "Lesson Blueprint")
    if len(blueprint_rows) != 8:
        raise SystemExit("Lesson Blueprint must contain exactly eight lesson rows.")
    not_approved = [str(row.get("Lesson")) for row in blueprint_rows if clean(row.get("Designer Status")) != "Approved"]
    if not_approved:
        raise SystemExit("Approve every Lesson Blueprint row before export. Pending lessons: " + ", ".join(not_approved))

    for row in blueprint_rows:
        number = int(row.get("Lesson"))
        steps = {}
        for suffix in ("2", "3", "4"):
            source = content_rows.get((number, suffix))
            requested = suffix == "2" or clean(row.get(f"Use .{suffix}?")) == "Yes"
            if not requested:
                steps[suffix] = None
                continue
            if not source:
                raise SystemExit(f"Lesson {number}.{suffix} is required by the approved blueprint but lacks approved Lesson Content.")
            approved_checks = check_rows.get((number, suffix), [])
            override = clean(source.get("Check Override Reason"))
            if not approved_checks and not override:
                raise SystemExit(f"Lesson {number}.{suffix} requires an approved check or an approved override reason.")
            graphic = None
            if clean(source.get("Graphic Path")):
                graphic = {
                    "path": clean(source.get("Graphic Path")),
                    "alt": clean(source.get("Graphic Alt Text")),
                    "caption": clean(source.get("Graphic Caption")),
                }
            resource_id = clean(source.get("Primary Resource ID"))
            steps[suffix] = {
                "type": "connect" if suffix == "2" else "learn",
                "title": clean(source.get("Step Title")),
                "focus": clean(source.get("Your Next Move")),
                "body": split_paragraphs(source.get("Student-Facing Instruction")),
                "resource_ids": [resource_id] if resource_id else [],
                "source_ids": split_values(source.get("Additional Source IDs")),
                "graphic": graphic,
                "check": approved_checks[0] if approved_checks else None,
                "check_override_reason": override,
                "estimated_minutes": source.get("Estimated Minutes") or 0,
            }

        application = app_rows.get(number)
        if not application:
            raise SystemExit(f"Lesson {number} requires an approved Application Activities row.")
        steps["application"] = {
            "type": "application",
            "title": clean(application.get("Activity Title")),
            "intro": clean(application.get("Purpose")),
            "completion_mode": clean(application.get("Completion Mode")).lower().replace("-", "_").replace(" ", "_"),
            "directions": split_paragraphs(application.get("Directions")),
            "program_scenarios": scenarios.get(number, []),
        }
        steps["assessment"] = {"type": "assessment", "title": f"Lesson {number} Assessment"}

        lessons.append({
            "number": number,
            "title": clean(row.get("Proposed Title")),
            "purpose": clean(row.get("Topic and Purpose")),
            "standards": split_values(row.get("Standard IDs")),
            "objectives": split_paragraphs(row.get("Draft Learning Objectives")),
            "vocabulary": parse_vocabulary(row.get("Vocabulary")),
            "estimated_minutes": int(row.get("Estimated Minutes") or 50),
            "steps": steps,
        })

    assessment_bank = []
    for row in sheet_rows(wb, "Assessment Bank"):
        if clean(row.get("Approval Status")) != "Approved":
            continue
        item = question(row)
        item.update({
            "lesson_number": int(row.get("Lesson")),
            "question_id": clean(row.get("Question ID")),
            "points": row.get("Points") or 0,
            "cognitive_level": clean(row.get("Cognitive Level")),
            "randomization_group": clean(row.get("Randomization Group")),
            "required": clean(row.get("Required")),
            "reflection_item": clean(row.get("Reflection Item")),
            "approval_status": "Approved",
            "google_forms_override": clean(row.get("Google Forms Override")),
            "canvas_override": clean(row.get("Canvas Override")),
            "conversion_notes": clean(row.get("Conversion Notes")),
        })
        for choice in item.pop("choices", []):
            item[f"option_{choice['id']}"] = choice["text"]
        item["correct_answers"] = ";".join(item.pop("correct", []))
        assessment_bank.append(item)

    data = {
        "course": {
            "title": clean(setup.get("Course title")) or "Course Module Title",
            "short_title": clean(setup.get("Short title")) or "Course Module",
            "description": clean(setup.get("Module purpose")),
            "audience": "Grades 11–12 career and technical education students",
            "module_type": "cte" if clean(setup.get("Module type")).lower() == "cte" else "general",
            "lesson_count": 8,
            "target_minutes_per_lesson": 50,
            "reading_level_target": "7–8",
            "feedback_form_url": clean(setup.get("Feedback form URL")),
            "mastery": {"threshold_percent": 80, "attempts": 3, "retain": "highest", "shuffle_questions": True, "shuffle_choices": True, "show_correct_answers": False, "release_results": "instructor"},
        },
        "standards": standards,
        "resources": resources,
        "lessons": lessons,
        "assessment_bank": assessment_bank,
    }
    OUTPUT.write_text(json.dumps(data, indent=2), encoding="utf-8")
    print(f"Exported approved workbook content to {OUTPUT}")


if __name__ == "__main__":
    main()
